"""Reading time series files (CSV and Excel).

The reader works out the file layout by itself and the user can correct every
decision:

- encoding        UTF-8 (with or without BOM), Windows-1252, UTF-16
- delimiter       semicolon, comma, tab or pipe
- decimal sign    "12,5" (German) or "12.5", also with thousands separators
- header          lines with a title or metadata above the column names are
                  skipped, a unit line below the column names is skipped
- Excel           .xlsx and .xlsm, any sheet, numbers stored as text

This module does not use Streamlit, so any frontend can call it.
"""

import csv
import datetime
import io
import math
import numbers
import re
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

CSV_EXTENSIONS = (".csv", ".txt", ".tsv")
EXCEL_EXTENSIONS = (".xlsx", ".xlsm")
FILE_TYPES = [ext.lstrip(".") for ext in CSV_EXTENSIONS + EXCEL_EXTENSIONS]

# tried in this order when a delimiter is detected (the first one wins a tie)
DELIMITER_CANDIDATES = ("\t", ";", ",", "|")

# cells that mean "no value"
MISSING_TOKENS = {"", "nan", "n/a", "na", "#n/a", "null", "none", "-", "--", "\u2014"}

# a column counts as dates when at least this share of its cells is a date,
# and as numbers when at least this share is a number (the others are reported)
MIN_SHARE_DATES = 0.9
MIN_SHARE_NUMBERS = 0.5


class TimeSeriesFileError(ValueError):
    """The file cannot be read. The message is written for the user."""


@dataclass(frozen=True)
class ImportOptions:
    """What the user (or the detection) decided. None means: detect it."""

    encoding: str = None
    delimiter: str = None
    decimal: str = None            # "." or ","
    header_line: int = None        # line of the column names, 0 = no header
    first_data_line: int = None    # first line with values
    sheet: str = None              # Excel only


@dataclass
class TimeAxis:
    column: str
    start: pd.Timestamp
    end: pd.Timestamp
    rows: int
    step_minutes: float
    sorted: bool
    duplicates: int
    gaps: int


@dataclass
class ImportedTable:
    df: pd.DataFrame
    kind: str                       # "csv" or "excel"
    used: ImportOptions             # the options that were really applied
    sheets: list = field(default_factory=list)
    messages: list = field(default_factory=list)    # (level, text) pairs
    row_lines: list = field(default_factory=list)   # file line of every row
    time: TimeAxis = None


# ---------------------------------------------------------------------------
# helpers for single cells
# ---------------------------------------------------------------------------

_NUMBER_RE = re.compile(
    r"^[+\-\u2212]?\s*(?:\d[\d\s.,'\u00a0\u202f]*|[.,]\d+)(?:[eE][+\-]?\d+)?$"
)
_DATE_RE = re.compile(
    r"^\s*(?:\d{4}[-/.]\d{1,2}[-/.]\d{1,2}|\d{1,2}[-/.]\d{1,2}[-/.]\d{2,4})"
    r"(?:[ T]+\d{1,2}:\d{2}(?::\d{2}(?:[.,]\d+)?)?)?"
    r"\s*(?:Z|[+\-]\d{2}:?\d{2})?\s*$"
)


def _is_blank(value):
    """No content at all (used to find the table in the sheet)."""

    if value is None:
        return True
    if isinstance(value, float) and math.isnan(value):
        return True
    return isinstance(value, str) and not value.strip()


def _is_missing(value):
    """Blank, or a text such as n/a that stands for 'no value'."""

    if _is_blank(value):
        return True
    return isinstance(value, str) and value.strip().lower() in MISSING_TOKENS


def _is_number(value):
    return isinstance(value, numbers.Number) and not isinstance(value, bool)


def _is_datetime_object(value):
    return isinstance(value, (datetime.datetime, datetime.date, pd.Timestamp))


def _looks_like_data(value):
    """True for numbers and dates, False for words (column names, titles)."""

    if _is_number(value) or _is_datetime_object(value):
        return True
    if isinstance(value, str):
        text = value.strip()
        return bool(_NUMBER_RE.match(text) or _DATE_RE.match(text))
    return False


# ---------------------------------------------------------------------------
# reading the raw cells
# ---------------------------------------------------------------------------

def _decode(content, encoding):
    if not content.strip():
        raise TimeSeriesFileError("The file is empty.")

    if encoding:
        try:
            return content.decode(encoding), encoding
        except (UnicodeDecodeError, LookupError):
            raise TimeSeriesFileError(
                f"The file cannot be read with the encoding '{encoding}'. "
                f"Choose 'Automatic' or another encoding."
            )

    if content.startswith((b"\xff\xfe", b"\xfe\xff")):
        return content.decode("utf-16"), "utf-16"

    for candidate in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(candidate), candidate
        except UnicodeDecodeError:
            continue

    return content.decode("latin-1"), "latin-1"


def _detect_delimiter(text):
    lines = [line for line in text.splitlines() if line.strip()][:300]

    best, best_share = None, 0.0

    for candidate in DELIMITER_CANDIDATES:
        counts = [len(row) for row in csv.reader(lines, delimiter=candidate)]
        modal, frequency = Counter(counts).most_common(1)[0]

        if modal < 2:
            continue

        share = frequency / len(counts)
        if share >= 0.5 and share > best_share:
            best, best_share = candidate, share

    return best or ","


def _csv_grid(content, options):
    text, encoding = _decode(content, options.encoding)
    delimiter = options.delimiter or _detect_delimiter(text)

    try:
        grid = list(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter))
    except csv.Error as error:
        raise TimeSeriesFileError(f"The file cannot be read as a table: {error}")

    return grid, encoding, delimiter


def _excel_grid(content, options):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        raise TimeSeriesFileError(
            "The Excel file cannot be opened. Is it damaged or password protected?"
        )

    try:
        sheets = list(workbook.sheetnames)

        def read(name):
            sheet = workbook[name]
            if hasattr(sheet, "reset_dimensions"):
                sheet.reset_dimensions()
            return [list(row) for row in sheet.iter_rows(values_only=True)]

        if options.sheet:
            if options.sheet not in sheets:
                raise TimeSeriesFileError(f"The sheet '{options.sheet}' does not exist.")
            return read(options.sheet), sheets, options.sheet

        # the sheet with the most numbers and dates is the data sheet; a sheet
        # with notes or a title page has none
        best = None
        for name in sheets:
            grid = read(name)
            filled = sum(not _is_blank(cell) for row in grid for cell in row)
            data = sum(_looks_like_data(cell) for row in grid for cell in row)

            if filled and (best is None or (data, filled) > best[0]):
                best = ((data, filled), grid, name)

        if best is None:
            raise TimeSeriesFileError("The Excel file contains no data.")

        return best[1], sheets, best[2]
    finally:
        workbook.close()


# ---------------------------------------------------------------------------
# finding the table inside the grid
# ---------------------------------------------------------------------------

def _pad(grid):
    width = max((len(row) for row in grid), default=0)
    return [list(row) + [None] * (width - len(row)) for row in grid], width


def _find_table(grid):
    """Detect the header line and the first data line (0-based indices).

    Returns (header_index or None, data_index).
    """

    counts = [sum(not _is_blank(cell) for cell in row) for row in grid]
    filled = [c for c in counts if c > 0]

    if not filled:
        raise TimeSeriesFileError("The file contains no data.")

    # the usual number of filled cells in a row is the width of the table
    frequency = Counter(filled)
    width = max(frequency, key=lambda c: (frequency[c], c))

    def share_of_data(index):
        cells = [c for c in grid[index] if not _is_blank(c)]
        return sum(_looks_like_data(c) for c in cells) / len(cells) if cells else 0.0

    candidates = [i for i, c in enumerate(counts) if c == width]
    first = next(
        (i for i in candidates if i + 1 < len(counts) and counts[i + 1] >= width - 1),
        candidates[0],
    )

    # column names: words, followed (maybe after a unit line) by numbers
    if share_of_data(first) < 0.5:
        for later in range(first + 1, min(first + 6, len(grid))):
            if counts[later] >= max(width - 1, 1) and share_of_data(later) >= 0.5:
                return first, later
        return first, first + 1

    # no column names in the first full row: look for a line with only words above it
    above = first - 1
    while above >= 0 and counts[above] == 0:
        above -= 1
    if above >= 0 and counts[above] >= max(width - 1, 1) and share_of_data(above) < 0.5:
        return above, first

    return None, first


def _unique_names(names):
    seen, result = Counter(), []
    for position, name in enumerate(names, start=1):
        text = "" if _is_blank(name) else str(name).strip()
        text = text or f"Column {position}"
        seen[text] += 1
        result.append(text if seen[text] == 1 else f"{text} ({seen[text]})")
    return result


# ---------------------------------------------------------------------------
# converting the columns
# ---------------------------------------------------------------------------

def _column_kind(values):
    """'empty', 'number', 'datetime' or 'text'."""

    present = [v for v in values if not _is_missing(v)]

    if not present:
        return "empty"

    if all(_is_number(v) for v in present):
        return "number"

    dates = sum(
        _is_datetime_object(v) or (isinstance(v, str) and bool(_DATE_RE.match(v.strip())))
        for v in present
    )
    if dates / len(present) >= MIN_SHARE_DATES:
        return "datetime"

    numbers_found = sum(
        _is_number(v) or (isinstance(v, str) and bool(_NUMBER_RE.match(v.strip())))
        for v in present
    )
    if numbers_found / len(present) >= MIN_SHARE_NUMBERS:
        return "number"

    return "text"


def _decimal_votes(text):
    """What a single number tells about the decimal sign.

    Returns ',', '.', 'ambiguous-comma', 'ambiguous-dot' or None.
    """

    s = re.sub(r"[\s\u00a0\u202f']", "", text.replace("\u2212", "-"))
    s = re.sub(r"[eE][+\-]?\d+$", "", s).lstrip("+-")

    has_comma, has_dot = "," in s, "." in s

    if has_comma and has_dot:
        return "," if s.rfind(",") > s.rfind(".") else "."

    for sign, ambiguous in ((",", "ambiguous-comma"), (".", "ambiguous-dot")):
        if sign not in s:
            continue
        if s.count(sign) > 1:                       # 1.234.567 or 1,234,567
            return "." if sign == "," else ","
        lead, after = s.split(sign)
        if len(after) != 3 or lead.startswith("0") or lead == "":
            return sign                              # 12,5 / 0,125 / ,5
        return ambiguous                             # 1,234 could be either

    return None


def _detect_decimal(strings, delimiter):
    """Decide the decimal sign of a whole file. Returns (sign, certain)."""

    votes = Counter(_decimal_votes(s) for s in strings)

    comma, dot = votes[","], votes["."]

    if comma and not dot:
        return ",", True
    if dot and not comma:
        return ".", True
    if comma and dot:
        return (",", False) if comma > dot else (".", False)

    # only ambiguous values such as 1,234 or 1.234
    if votes["ambiguous-comma"] or votes["ambiguous-dot"]:
        # in a semicolon or tab separated file the decimal sign is usually a comma
        return (",", False) if delimiter in (";", "\t") else (".", False)

    return ".", True


def _parse_number(value, decimal):
    if _is_number(value):
        return float(value)
    if _is_missing(value):
        return math.nan

    s = str(value).strip().replace("\u2212", "-")
    s = re.sub(r"[\s\u00a0\u202f']", "", s)

    if decimal == ",":
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", "")

    try:
        return float(s)
    except ValueError:
        return math.nan


def _day_first(strings):
    """German dates (31.12.2021) are day first; 12/31/2021 is month first."""

    for text in strings:
        if re.match(r"^\s*\d{4}[-/.]", text):
            return False                              # ISO: year first
        parts = re.split(r"[-/.\s]", text.strip())
        if len(parts) >= 2 and parts[0].isdigit() and parts[1].isdigit():
            if int(parts[0]) > 12:
                return True
            if int(parts[1]) > 12:
                return False
    return True


def _parse_datetimes(values):
    texts = [v.strip() for v in values if isinstance(v, str) and not _is_missing(v)]
    day_first = _day_first(texts[:2000])

    series = pd.Series(
        [None if _is_missing(v) else v for v in values], dtype=object
    )

    try:
        parsed = pd.to_datetime(series, dayfirst=day_first, errors="coerce", format="mixed")
    except (ValueError, TypeError):
        parsed = pd.to_datetime(series, dayfirst=day_first, errors="coerce")

    if getattr(parsed.dt, "tz", None) is not None:
        parsed = parsed.dt.tz_localize(None)         # keep the wall clock time

    return parsed


def _time_axis(df):
    """Describe the first date column (step, gaps, duplicates)."""

    for column in df.columns:
        if not pd.api.types.is_datetime64_any_dtype(df[column]):
            continue

        stamps = df[column].dropna()
        if len(stamps) < 2:
            return None

        steps = stamps.diff().dropna()
        step = steps.mode().iloc[0]

        return TimeAxis(
            column=column,
            start=stamps.iloc[0],
            end=stamps.iloc[-1],
            rows=len(stamps),
            step_minutes=step.total_seconds() / 60,
            sorted=bool(stamps.is_monotonic_increasing),
            duplicates=int(stamps.duplicated().sum()),
            gaps=int((steps > step).sum()),
        )

    return None


def _describe_time(axis):
    messages = []

    def stamp(value):
        return value.strftime("%d.%m.%Y %H:%M")

    step = axis.step_minutes
    if step < 60:
        step_text = f"{step:g} minutes"
    else:
        step_text = "1 hour" if step == 60 else f"{step / 60:g} hours"

    messages.append((
        "info",
        f"Time column '{axis.column}': {stamp(axis.start)} to {stamp(axis.end)}, "
        f"{axis.rows} rows, time step {step_text}.",
    ))

    if step != 60:
        messages.append((
            "warning",
            f"The time step of the file is {step_text}, but the model works with "
            f"hourly values. Convert the data to hourly values first.",
        ))
    if not axis.sorted:
        messages.append((
            "warning",
            "The timestamps are not in ascending order. The model uses the rows "
            "in the order of the file.",
        ))
    if axis.duplicates:
        messages.append((
            "warning",
            f"{axis.duplicates} timestamps appear twice. This is typical for the "
            f"end of daylight saving time (the hour from 02:00 to 03:00 is "
            f"recorded twice).",
        ))
    if axis.gaps:
        messages.append((
            "warning",
            f"{axis.gaps} gaps in the timestamps. This is typical for the start "
            f"of daylight saving time (one hour is missing) or for missing data.",
        ))

    return messages


# ---------------------------------------------------------------------------
# the public functions
# ---------------------------------------------------------------------------

def read_table(content, filename, options=None):
    """Read a CSV or Excel file into a table with numbers and dates converted.

    content   the bytes of the file
    filename  used to tell CSV from Excel (by the extension)
    options   ImportOptions; everything left at None is detected

    Raises TimeSeriesFileError (with a message for the user) if the file cannot
    be read.
    """

    options = options or ImportOptions()
    extension = Path(filename).suffix.lower()

    sheets = []

    if extension in EXCEL_EXTENSIONS:
        kind = "excel"
        grid, sheets, sheet = _excel_grid(content, options)
        encoding = delimiter = None
    elif extension in CSV_EXTENSIONS:
        kind = "csv"
        grid, encoding, delimiter = _csv_grid(content, options)
        sheet = None
    elif extension == ".xls":
        raise TimeSeriesFileError(
            "The old Excel format (.xls) is not supported. Open the file in "
            "Excel and save it as .xlsx, or export it as CSV."
        )
    else:
        raise TimeSeriesFileError(
            f"The file type '{extension or 'without extension'}' is not supported. "
            f"Use CSV, TXT, TSV or Excel (.xlsx)."
        )

    grid, width = _pad(grid)

    if not any(not _is_blank(cell) for row in grid for cell in row):
        raise TimeSeriesFileError("The file contains no data.")

    # --- where is the table? (line numbers are 1-based, like in an editor)
    header_index, data_index = _find_table(grid)

    if options.header_line is not None:
        header_index = None if options.header_line == 0 else options.header_line - 1
        if options.first_data_line is None:
            data_index = (header_index + 1) if header_index is not None else data_index

    if options.first_data_line is not None:
        data_index = max(options.first_data_line - 1, 0)
        if header_index is not None and header_index >= data_index:
            header_index = None

    rows = [
        (index, row) for index, row in enumerate(grid)
        if index >= data_index and any(not _is_blank(cell) for cell in row)
    ]

    if not rows:
        raise TimeSeriesFileError(
            "No data rows found below the column names. Check the line numbers "
            "in the import settings."
        )

    # drop columns that are empty everywhere (also the header)
    keep = [
        col for col in range(width)
        if any(not _is_blank(row[col]) for _, row in rows)
        or (header_index is not None and not _is_blank(grid[header_index][col]))
    ]

    names = _unique_names(
        [grid[header_index][col] for col in keep] if header_index is not None
        else [None] * len(keep)
    )
    columns = {name: [row[col] for _, row in rows] for name, col in zip(names, keep)}
    row_lines = [index + 1 for index, _ in rows]

    # --- what is in the columns?
    kinds = {name: _column_kind(values) for name, values in columns.items()}

    text_for_decimal = [
        str(v).strip()
        for name, values in columns.items() if kinds[name] == "number"
        for v in values if isinstance(v, str) and not _is_missing(v)
        and not _DATE_RE.match(v.strip())
    ]

    if options.decimal:
        decimal, certain = options.decimal, True
    elif text_for_decimal:
        decimal, certain = _detect_decimal(text_for_decimal, delimiter)
    else:
        decimal, certain = ".", True

    messages = []
    data = {}

    for name, values in columns.items():
        kind_of_column = kinds[name]

        if kind_of_column == "datetime":
            parsed = _parse_datetimes(values)
            bad = [i for i, v in enumerate(values) if not _is_missing(v) and pd.isna(parsed.iloc[i])]
            data[name] = parsed
            if bad:
                messages.append((
                    "warning",
                    f"Column '{name}': {len(bad)} dates could not be read "
                    f"(first in line {row_lines[bad[0]]}).",
                ))

        elif kind_of_column == "number":
            converted = [_parse_number(v, decimal) for v in values]
            data[name] = pd.Series(converted, dtype=float)

            not_numbers = [
                i for i, (v, c) in enumerate(zip(values, converted))
                if math.isnan(c) and not _is_missing(v)
            ]
            empty = [
                i for i, v in enumerate(values) if _is_missing(v)
            ]
            if not_numbers:
                messages.append((
                    "warning",
                    f"Column '{name}': {len(not_numbers)} values are not numbers "
                    f"(first in line {row_lines[not_numbers[0]]}: "
                    f"'{values[not_numbers[0]]}').",
                ))
            if empty:
                messages.append((
                    "warning",
                    f"Column '{name}': {len(empty)} empty values "
                    f"(first in line {row_lines[empty[0]]}).",
                ))

        else:
            data[name] = pd.Series(
                [None if _is_blank(v) else str(v).strip() for v in values], dtype=object
            )

    df = pd.DataFrame(data)

    used = ImportOptions(
        encoding=encoding,
        delimiter=delimiter,
        decimal=decimal,
        header_line=0 if header_index is None else header_index + 1,
        first_data_line=data_index + 1,
        sheet=sheet,
    )

    # --- plain messages about what was found
    if any("," in t or "." in t for t in text_for_decimal):
        sign = "comma (12,5)" if decimal == "," else "point (12.5)"
        messages.insert(0, (
            "info" if certain else "warning",
            f"Decimal sign: {sign}" + ("" if certain else " - assumed, please check the preview."),
        ))

    table = ImportedTable(
        df=df, kind=kind, used=used, sheets=sheets,
        messages=messages, row_lines=row_lines,
    )

    table.time = _time_axis(df)
    if table.time is not None:
        table.messages = _describe_time(table.time) + table.messages
    else:
        table.messages.insert(0, (
            "info", "No date column found. The rows are used in the order of the file."
        ))

    return table


def series_from_column(table, column):
    """The values of one column as a list of numbers (empty cells stay NaN)."""

    if column not in table.df.columns:
        raise TimeSeriesFileError(f"The column '{column}' does not exist.")

    series = table.df[column]

    if not pd.api.types.is_numeric_dtype(series):
        raise TimeSeriesFileError(
            f"The column '{column}' does not contain numbers. Choose another column."
        )

    return series.tolist()